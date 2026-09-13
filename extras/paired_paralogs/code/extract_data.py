from pathlib import Path
import csv,json,math
import openpyxl
R=Path(__file__).resolve().parents[1];s=openpyxl.load_workbook(R/'sources/Data.xlsx',data_only=True).active
cols={'GlnB_total':9,'GlnB_unmodified':12,'GlnB_UMP':15,'GlnK_total':18,'GlnK_unmodified':21,'GlnK_UMP':24,'GS_total':27,'GS_unmodified':30,'GS_AMP':33,'aKG':36,'glutamine':39}
(R/'results').mkdir(exist_ok=True)
rows=[]
for strain,ran in [('WT',range(8,20)),('delta_glnB',range(25,37)),('delta_glnK',range(42,54))]:
 for ri in ran:
  for name,c in cols.items():
   val=s.cell(ri,c).value
   if not isinstance(val,(float,int)):continue
   rows.append(dict(strain=strain,source_row=ri,source_cell=s.cell(ri,c).coordinate,phase=s.cell(ri,2).value,time_min=s.cell(ri,3).value,time_SE=s.cell(ri,4).value,observable=name,mean_copy_cell=val,SE_copy_cell=s.cell(ri,c+1).value,n=s.cell(ri,c+2).value))
with (R/'results/observations.csv').open('w') as f:
 w=csv.DictWriter(f,fieldnames=rows[0]);w.writeheader();w.writerows(rows)
# The published conversion factor, not a new cell-volume estimate.
conv=.00161
for strain in ['WT','delta_glnB']:
 print(strain)
 for ri in sorted({x['source_row'] for x in rows if x['strain']==strain}):
  d={x['observable']:x for x in rows if x['source_row']==ri}
  if not all(k in d for k in ['GlnK_total','GlnK_UMP','aKG']):continue
  T,U,A=[d[k]['mean_copy_cell'] for k in ['GlnK_total','GlnK_UMP','aKG']];a=A*conv/1000;fraction=U/T;accessible=.4+.6*a**3/(3**3+a**3)
  print(d['aKG']['time_min'],'theta',round(fraction,3),'accessible',round(accessible,3),'gap',round(fraction-accessible,3))
(R/'results/extraction.json').write_text(json.dumps({'source_doi':'10.6084/m9.figshare.4880003','article_doi':'10.1016/j.bpj.2017.04.012','source_sheet':'Sheet1','records':len(rows),'units':'source copy/cell; conversion 0.00161 micromolar per copy/cell follows original model','uncertainty':'source SE and n retained; no replicate vectors or cross-observable covariance provided','note':'Workbook cell B2 prints 1.61 M per 1000 molecules, inconsistent with 1 fl. Published model and physical scale support 1.61 micromolar; preserve raw text and document conversion explicitly.'},indent=2))
